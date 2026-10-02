from __future__ import annotations

from io import BytesIO
from typing import BinaryIO

from openpyxl import Workbook, load_workbook

from .data import TABLES
from .storage import get_extra_fields, get_visible_fields


def build_excel_template(table_key: str, labels: list[str]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Carga"
    sheet.append(labels)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def parse_excel_records(
    table_key: str,
    file_stream: BinaryIO,
    archived: bool = False,
) -> list[dict[str, str]]:
    workbook = load_workbook(file_stream, read_only=True, data_only=True)
    sheet = workbook.active
    visible_fields = get_visible_fields(table_key)
    records: list[dict[str, str]] = []

    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index == 1 and looks_like_header(table_key, row):
            continue
        values = [format_cell(value) for value in row[: len(visible_fields)]]
        values.extend([""] * (len(visible_fields) - len(values)))
        if not any(values):
            continue
        record = dict(zip(visible_fields, values))
        for field in get_extra_fields(table_key):
            record[field] = "1" if archived else "0"
        records.append(record)

    return records


def looks_like_header(table_key: str, row: tuple[object, ...]) -> bool:
    expected = [
        normalize_header(label)
        for _field, label in TABLES[table_key]["columns"]
    ]
    received = [normalize_header(value) for value in row[: len(expected)]]
    return received == expected


def normalize_header(value: object) -> str:
    return str(value or "").strip().lower()


def format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
